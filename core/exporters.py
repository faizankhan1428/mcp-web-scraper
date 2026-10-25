import os
from datetime import datetime
from fpdf import FPDF
from docx import Document

def save_as_markdown(content: str, url: str, folder: str, filename: str) -> str:
    if not filename.endswith(".txt") and not filename.endswith(".md"):
        filename += ".md"
    file_path = os.path.join(folder, filename)
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(f"# Source URL: {url}\n")
        f.write(f"# Scraped Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
        f.write(content[:5000])
    return file_path

def save_as_pdf(content: str, url: str, folder: str, filename: str) -> str:
    if not filename.endswith(".pdf"):
        filename += ".pdf"
    file_path = os.path.join(folder, filename)
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, f"Web Scrape Report - {datetime.now().strftime('%Y-%m-%d')}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "I", 10)
    pdf.cell(0, 6, f"Source URL: {url}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    clean_text = content[:4000].encode('latin-1', 'replace').decode('latin-1')
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, clean_text)
    pdf.output(file_path)
    return file_path

def save_as_docx(content: str, url: str, folder: str, filename: str) -> str:
    if not filename.endswith(".docx"):
        filename += ".docx"
    file_path = os.path.join(folder, filename)
    
    doc = Document()
    doc.add_heading("Web Scrape Summary Report", level=1)
    doc.add_paragraph(f"Source URL: {url}")
    doc.add_paragraph(f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_heading("Extracted Article Content", level=2)

    paragraphs = [p.strip() for p in content[:5000].split("\n") if p.strip()]
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(file_path)
    return file_path