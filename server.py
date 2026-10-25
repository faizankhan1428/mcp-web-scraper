from mcp.server.mcpserver import MCPServer
import httpx
from bs4 import BeautifulSoup
import os
import csv
import sqlite3
from datetime import datetime
from fpdf import FPDF
from docx import Document

# Import AI Summary function from our core folder
from core.ai_summary import generate_smart_summary
# Import Smart Full-Text Search Engine functions from core folder
from core.search_engine import init_fts_db, index_document_for_search, smart_search

# Initialize the MCP server
mcp = MCPServer("Advanced Web Scraper MCP")

# Define folder path for saving files locally
SAVE_FOLDER = r"D:\Faizan Data\mcp-web-scraper\scraped_outputs"
os.makedirs(SAVE_FOLDER, exist_ok=True)

# Database path
DB_PATH = os.path.join(SAVE_FOLDER, "scraped_database.db")

def init_db():
    """Initializes the SQLite database tables for scraping history and FTS smart search."""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    # Standard log table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS scrapes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            timestamp TEXT,
            url TEXT,
            title TEXT,
            file_path TEXT,
            file_type TEXT
        )
    """)
    conn.commit()
    conn.close()
    
    # Initialize FTS5 Smart Search Table
    init_fts_db(DB_PATH)

# Initialize database on startup
init_db()


def log_to_database(url: str, title: str, file_path: str, file_type: str, content: str = ""):
    """Helper function to log scrape entry into SQLite database and index for FTS search."""
    try:
        timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # 1. Log to standard history table
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO scrapes (timestamp, url, title, file_path, file_type) VALUES (?, ?, ?, ?, ?)",
            (timestamp_str, url, title, file_path, file_type)
        )
        conn.commit()
        conn.close()

        # 2. Index into Smart FTS5 Search Engine
        if content:
            index_document_for_search(DB_PATH, title, url, content, file_path, file_type, timestamp_str)

    except Exception as e:
        print(f"DB Error: {e}")


@mcp.tool()
def scrape_website(url: str) -> str:
    """
    Fetches a website URL and extracts clean text content.
    """
    try:
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        response = httpx.get(url, headers=headers, timeout=10.0, follow_redirects=True)
        
        if response.status_code != 200:
            return f"Error: Failed to fetch website. Status code: {response.status_code}"

        soup = BeautifulSoup(response.text, "html.parser")

        for script in soup(["script", "style", "nav", "footer", "header"]):
            script.decompose()

        text = soup.get_text(separator="\n", strip=True)
        return text

    except Exception as e:
        return f"An error occurred: {str(e)}"


@mcp.tool()
def scrape_and_save_file(url: str, filename: str, page_title: str = "Untitled Page") -> str:
    """
    Scrapes a website, generates AI summary, saves content to a Markdown/Text file, and logs/indexes it into DB.
    """
    try:
        content = scrape_website(url)
        if content.startswith("Error"):
            return content

        # Generate AI Summary for Phase 5
        ai_summary = generate_smart_summary(content)

        if not filename.endswith(".txt") and not filename.endswith(".md"):
            filename += ".md"
            
        file_path = os.path.join(SAVE_FOLDER, filename)

        with open(file_path, "w", encoding="utf-8") as f:
            f.write(f"# Source URL: {url}\n")
            f.write(f"# Scraped Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
            f.write(ai_summary)  # Injecting AI Summary at the top
            f.write(content[:5000])

        log_to_database(url, page_title, file_path, "Markdown/Text", content)
        return f"Success! Saved file with AI Summary and indexed in Smart Search DB at: {file_path}"

    except Exception as e:
        return f"Failed to save file: {str(e)}"


@mcp.tool()
def scrape_and_save_pdf(url: str, filename: str, page_title: str = "PDF Report") -> str:
    """
    Scrapes a website URL, adds AI summary, generates a formatted PDF document, and logs/indexes it into DB.
    """
    try:
        content = scrape_website(url)
        if content.startswith("Error"):
            return content

        # Generate AI Summary for Phase 5
        ai_summary = generate_smart_summary(content)

        if not filename.endswith(".pdf"):
            filename += ".pdf"
            
        file_path = os.path.join(SAVE_FOLDER, filename)

        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, f"Web Scrape Report & AI Insights - {datetime.now().strftime('%Y-%m-%d')}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 6, f"Source URL: {url}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

        # Encode summary and content safely for FPDF
        combined_text = ai_summary + "\n\n" + content[:4000]
        clean_text = combined_text.encode('latin-1', 'replace').decode('latin-1')
        
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, clean_text)

        pdf.output(file_path)
        log_to_database(url, page_title, file_path, "PDF", content)
        return f"Success! Formatted PDF with AI Summary generated and indexed in Smart Search DB at: {file_path}"

    except Exception as e:
        return f"Failed to create PDF file: {str(e)}"


@mcp.tool()
def scrape_and_save_docx(url: str, filename: str, page_title: str = "Word Document") -> str:
    """
    Scrapes a website URL, includes AI summary, generates a Word document (.docx), and logs/indexes it into DB.
    """
    try:
        content = scrape_website(url)
        if content.startswith("Error"):
            return content

        # Generate AI Summary for Phase 5
        ai_summary = generate_smart_summary(content)

        if not filename.endswith(".docx"):
            filename += ".docx"
            
        file_path = os.path.join(SAVE_FOLDER, filename)

        doc = Document()
        doc.add_heading("Web Scrape Summary & AI Insights Report", level=1)
        doc.add_paragraph(f"Source URL: {url}")
        doc.add_paragraph(f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        
        # Add AI Summary Section in Word
        doc.add_heading("AI Executive Summary", level=2)
        for line in ai_summary.split("\n"):
            if line.strip():
                doc.add_paragraph(line)

        doc.add_heading("Full Extracted Article Content", level=2)
        paragraphs = [p.strip() for p in content[:5000].split("\n") if p.strip()]
        for p in paragraphs:
            doc.add_paragraph(p)

        doc.save(file_path)
        log_to_database(url, page_title, file_path, "Word", content)
        return f"Success! Word Document with AI Summary generated and indexed in Smart Search DB at: {file_path}"

    except Exception as e:
        return f"Failed to create Word document: {str(e)}"


@mcp.tool()
def search_scraped_database(keyword: str) -> str:
    """
    Performs intelligent full-text relevance search across all past scraped documents using FTS5.
    """
    return smart_search(DB_PATH, keyword)


@mcp.tool()
def log_scrape_to_csv(url: str, page_title: str) -> str:
    """
    Logs the scraped history into a CSV spreadsheet file.
    """
    try:
        csv_file_path = os.path.join(SAVE_FOLDER, "scraped_history.csv")
        file_exists = os.path.isfile(csv_file_path)

        with open(csv_file_path, mode="a", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            if not file_exists:
                writer.writerow(["Timestamp", "URL", "Page Title"])
            writer.writerow([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), url, page_title])

        return f"Logged successfully into local spreadsheet: {csv_file_path}"

    except Exception as e:
        return f"Failed to log to CSV: {str(e)}"


if __name__ == "__main__":
    mcp.run()